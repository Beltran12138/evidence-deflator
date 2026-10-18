"""Export the Macau deck to PDF and PPTX, one full-slide image per page, for organisers who need a file.

    python -m http.server 8765 --directory docs      # in another terminal
    python tools/export_deck.py OUT_DIR [--slides 0-7] [--base http://127.0.0.1:8765/macau2.html]

Each slide is opened with ?static=1 (animations off, final state) and captured by headless Chrome at
1920x1080. The HTML stays the only source to edit; re-run this after every change. The PPTX holds images,
so it opens the same everywhere but its text cannot be edited in PowerPoint.
"""
import argparse
import os
import subprocess
import tempfile

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"


def shoot(url, path, profile):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={profile}",
                    "--force-device-scale-factor=1", "--window-size=1920,1080", "--virtual-time-budget=9000",
                    f"--screenshot={path}", url], check=True, capture_output=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--slides", default="0-7")
    ap.add_argument("--base", default="http://127.0.0.1:8765/macau2.html")
    ap.add_argument("--name", default="macau-deck")
    a = ap.parse_args()
    lo, hi = (int(x) for x in a.slides.split("-"))
    os.makedirs(a.out, exist_ok=True)
    profile = tempfile.mkdtemp(prefix="deckshot-")
    pngs = []
    for i in range(lo, hi + 1):
        path = os.path.join(a.out, f"{a.name}-{i:02d}.png")
        shoot(f"{a.base}?s={i}&static=1", path, profile)
        pngs.append(path)
        print("slide", i, Image.open(path).size)
    imgs = [Image.open(p).convert("RGB") for p in pngs]
    pdf = os.path.join(a.out, f"{a.name}.pdf")
    imgs[0].save(pdf, save_all=True, append_images=imgs[1:], resolution=144)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    for p in pngs:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(p, 0, 0, width=prs.slide_width, height=prs.slide_height)
    pptx = os.path.join(a.out, f"{a.name}.pptx")
    prs.save(pptx)
    print("wrote", pdf, pptx)


if __name__ == "__main__":
    main()
